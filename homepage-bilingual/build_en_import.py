# -*- coding: utf-8 -*-
"""產生「一次性匯入英文 Meta」WPCode snippet（A：頁面、Blog、分類；B：產品）及審閱 Excel"""
import os
from en_meta_data import all_rows, kw_of, name_status
from en_body_status import body_status, RAW as BODY_RAW
from en_urls import url_of
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
HERE = os.path.dirname(os.path.abspath(__file__))
q = lambda s: "'" + s.replace('\\', '\\\\').replace("'", "\\'") + "'"
R = all_rows()

def snippet(name, flag, rows):
    posts = [r for r in R if r in rows and r[0] in ('page', 'post', 'product')]
    terms = [r for r in rows if r[0] == 'term']
    spec = [r for r in rows if r[0] == 'special']
    L = [f"/**\n * Kitchen AO – 英文 Meta 一次性匯入 {name}（2026-10-28）",
         " * 用法：WPCode › PHP Snippet › Auto Insert › Admin Only（或 Run Everywhere）→ Active → Save",
         " *       儲存後打開任何後台頁面，頂部會顯示「已匯入 N 項」→ 之後把此 snippet 轉為 Inactive",
         " * 只會填入空白的 English SEO 方格，已填的內容不會被覆蓋。需要「中英文版通用工具 v2」已啟用。\n */",
         "add_action( 'admin_init', function () {",
         f"\tif ( get_option( '{flag}' ) || ! current_user_can( 'manage_options' ) ) return;",
         "\t$n = 0;"]
    if posts:
        L.append("\t$posts = array(")
        for r in posts: L.append(f"\t\t{r[1]} => array( {q(r[3])}, {q(r[4])} ), // {r[2][:24]}")
        L.append("\t);")
        L.append("\tforeach ( $posts as $id => $v ) {\n\t\tif ( ! get_post( $id ) ) continue;\n\t\tif ( get_post_meta( $id, '_ao_en_title', true ) === '' ) { update_post_meta( $id, '_ao_en_title', $v[0] ); $n++; }\n\t\tif ( get_post_meta( $id, '_ao_en_desc', true ) === '' ) { update_post_meta( $id, '_ao_en_desc', $v[1] ); }\n\t}")
    if terms:
        L.append("\t$terms = array(")
        for r in terms: L.append(f"\t\t{r[1]} => array( {q(r[3])}, {q(r[4])} ), // {r[2][:24]}")
        L.append("\t);")
        L.append("\tforeach ( $terms as $id => $v ) {\n\t\tif ( ! term_exists( (int) $id ) ) continue;\n\t\tif ( get_term_meta( $id, '_ao_en_title', true ) === '' ) { update_term_meta( $id, '_ao_en_title', $v[0] ); $n++; }\n\t\tif ( get_term_meta( $id, '_ao_en_desc', true ) === '' ) { update_term_meta( $id, '_ao_en_desc', $v[1] ); }\n\t}")
    for r in spec:
        idexpr = "(int) get_option( 'page_for_posts' )" if r[1] == 'blog' else "( function_exists( 'wc_get_page_id' ) ? (int) wc_get_page_id( 'shop' ) : 0 )"
        L.append(f"\t$sid = {idexpr}; // {r[2]}\n\tif ( $sid && get_post_meta( $sid, '_ao_en_title', true ) === '' ) {{ update_post_meta( $sid, '_ao_en_title', {q(r[3])} ); update_post_meta( $sid, '_ao_en_desc', {q(r[4])} ); $n++; }}")
    L += [f"\tupdate_option( '{flag}', $n . ' 項 @ ' . current_time( 'mysql' ), false );",
          "\tset_transient( 'ao_en_import_msg', '" + name + "：已匯入 ' . $n . ' 項英文 Meta（已有內容的方格沒有改動）。請把此 snippet 轉為 Inactive。', 300 );",
          "} );",
          "add_action( 'admin_notices', function () {",
          "\t$m = get_transient( 'ao_en_import_msg' );",
          "\tif ( $m ) { echo '<div class=\"notice notice-success\"><p>' . esc_html( $m ) . '</p></div>'; delete_transient( 'ao_en_import_msg' ); }",
          "} );"]
    return '\n'.join(L) + '\n'

A = [r for r in R if r[0] in ('page', 'post', 'term', 'special')]
Bp = [r for r in R if r[0] == 'product']
open(os.path.join(HERE, 'en-meta-import-A-pages-blog-categories.txt'), 'w', encoding='utf-8').write(snippet('A（頁面、Blog、分類）', 'ao_en_import_a_20261028', A))
open(os.path.join(HERE, 'en-meta-import-B-products.txt'), 'w', encoding='utf-8').write(snippet('B（產品）', 'ao_en_import_b_20261028', Bp))

wb = Workbook(); ws = wb.active; ws.title = '全網英文 Meta'
TY = {'page': '頁面', 'post': 'Blog', 'product': '產品', 'term': '產品分類', 'special': '頁面（商店／Blog 列表）'}
ws.append(['類型', 'ID', '中文頁面', '中文網址', '英文網址（/en/）', 'English Title', '字元', 'English Meta Description', '字元', '主要英文關鍵字（月搜／GSC 曝光）', '狀態（網站英文產品名／英文版）', '匯入檔', '英文版內文（10-30）', '/en/ 剩餘中文', '頁面建立方式', '建議翻譯做法', '優先'])
for r in R:
    zu, eu = url_of(r[0], r[1])
    _bid = {'shop': 7, 'blog': 18559}.get(r[1], r[1]) if r[0] in ('page', 'special') else None
    _bs = body_status(_bid) if _bid else None
    ws.append([TY[r[0]], r[1], r[2], zu, eu, r[3], len(r[3]), r[4], len(r[4]), kw_of(r[0], r[1]), name_status(r[0], r[1]), 'B' if r[0] == 'product' else 'A'] + (list(_bs) if _bs else ['', '', '', '', '']))
for c in ws[1]: c.font = Font(name='Arial', bold=True, color='F7CE46'); c.fill = PatternFill('solid', fgColor='1A1A1A')
for row in ws.iter_rows(min_row=2):
    for c in row: c.font = Font(name='Arial'); c.alignment = Alignment(wrap_text=True, vertical='top')
for col, w in zip('ABCDEFGHIJKLMNOPQ', [12, 8, 26, 40, 40, 48, 6, 70, 6, 40, 40, 8, 22, 10, 26, 46, 22]): ws.column_dimensions[col].width = w
ws.freeze_panes = 'F2'; ws.auto_filter.ref = f'A1:Q{ws.max_row}'
Y = PatternFill('solid', fgColor='FFF2A8')
for row in ws.iter_rows(min_row=2):
    for c in (row[3], row[4]):
        if c.value: c.hyperlink = c.value; c.font = Font(name='Arial', color='0563C1', underline='single')
    if row[9].value:
        for c in row: c.fill = Y
    if row[10].value and str(row[10].value).startswith('英文版已停用'):
        for c in row: c.fill = PatternFill('solid', fgColor='E7E6E6'); c.font = Font(name='Arial', color='808080')
    elif row[10].value and not str(row[10].value).startswith('沿用'):
        row[10].fill = PatternFill('solid', fgColor='FAD4D4')
BODY_FILL = {'未': 'F8CBAD', '部': 'FFE699', '大': 'FFF2CC', '已': 'E2EFDA', '不': 'EDEDED'}
for row in ws.iter_rows(min_row=2):
    v = row[12].value
    if v: row[12].fill = PatternFill('solid', fgColor=BODY_FILL.get(v[0], 'FFFFFF'))
for c in ws[1][12:17]: c.font = Font(name='Arial', bold=True, color='F7CE46'); c.fill = PatternFill('solid', fgColor='7F3F00')
# 新分頁：英文內文翻譯清單
tr = wb.create_sheet('英文內文翻譯清單', 1)
tr.append(['優先', 'ID', '頁面', '中文網址', '英文網址', '英文版內文', '/en/ 剩餘中文', '中文版中文字數', '頁面建立方式', '建議翻譯做法', '完成日期'])
D = 'https://aoaodelivery.com'
rows_tr = []
for i, v in BODY_RAW.items():
    st = body_status(i)
    rows_tr.append([st[4], i, v[0], D + v[1], D + '/en' + v[1], st[0], st[1], v[4], st[2], st[3], ''])
PO = lambda p: {'P1': 1, 'P2': 2, 'P3': 3, 'P4': 4}.get(p[:2], 9)
SO = lambda s: {'未': 0, '部': 1, '大': 2, '已': 3, '不': 4}.get(s[0], 5)
for r in sorted(rows_tr, key=lambda r: (SO(r[5]) if r[0] == '—' else 0, PO(r[0]), -r[7])): tr.append(r)
for c in tr[1]: c.font = Font(name='Arial', bold=True, color='F7CE46'); c.fill = PatternFill('solid', fgColor='1A1A1A')
for row in tr.iter_rows(min_row=2):
    for c in row: c.font = Font(name='Arial'); c.alignment = Alignment(wrap_text=True, vertical='top')
    row[5].fill = PatternFill('solid', fgColor=BODY_FILL.get(row[5].value[0], 'FFFFFF'))
    for c in (row[3], row[4]): c.hyperlink = c.value; c.font = Font(name='Arial', color='0563C1', underline='single')
for col, w in zip('ABCDEFGHIJK', [24, 8, 20, 36, 40, 22, 10, 10, 30, 52, 12]): tr.column_dimensions[col].width = w
tr.freeze_panes = 'C2'; tr.auto_filter.ref = f'A1:K{tr.max_row}'
lg = wb.create_sheet('更新紀錄')
for line in [
    ('日期', '更新內容'),
    ('2026-10-28', 'A 匯入：頁面、Blog、產品分類英文 Meta（103 項）'),
    ('2026-10-28', 'B 匯入：產品英文 Meta（155 項）'),
    ('2026-10-29', 'B2：產品英文 Title 加「| Kitchen AO」（128 項；超過 60 字元的不加）'),
    ('2026-10-29', 'B3：8 個產品 Title／Description 跟網站手動修改的英文名（6063、6072、6077、6078、6274、6409、14306、19579）'),
    ('2026-10-29', '舊 EN SEO snippet 停用；全站英文 Meta 由「中英文版通用工具 v2」English SEO 方格提供'),
    ('2026-10-29', 'Blog 只保留中文版：TranslatePress 排除 /blog/、/blog/*；Rank Math 301 ^en/blog(/.*)?$ → /blog$1（灰色列）'),
    ('2026-10-30', '新增「英文內文翻譯清單」分頁及主表 M–Q 欄：檢查全站 56 頁 /en/ 內文是否仍是中文，並列出建議翻譯做法及優先次序'),
    ('2026-10-30', '網站已補 4 個英文產品名：10388、14342、20217、20225（與 Meta 一致，毋須重新匯入）；全部產品已有英文名'),
]:
    lg.append(list(line))
for c in lg[1]: c.font = Font(name='Arial', bold=True, color='F7CE46'); c.fill = PatternFill('solid', fgColor='1A1A1A')
lg.column_dimensions['A'].width = 12; lg.column_dimensions['B'].width = 110
wb.save(os.path.join(HERE, '全網英文Meta_2026-10-28.xlsx'))
print(len(A), len(Bp))
